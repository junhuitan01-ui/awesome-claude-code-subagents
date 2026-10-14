from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR
from common import OUT_DIR, parse_slides

BODY_PT = 22
TITLE_PT = 30


def main():
    slides = parse_slides()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for idx, s in enumerate(slides):
        if idx == 0:
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            title, sub = sl.shapes.title, sl.placeholders[1]
            title.left, title.top, title.width, title.height = Inches(0.9), Inches(1.5), Inches(11.5), Inches(1.9)
            sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(3.7), Inches(11.5), Inches(2.6)
            title.text_frame.text = s["title"]
            title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
            tf = sub.text_frame
            tf.word_wrap = True
            for k, b in enumerate(s["bullets"]):
                par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                par.text = b
                par.runs[0].font.size = Pt(24)
        else:
            sl = prs.slides.add_slide(prs.slide_layouts[1])
            title, body = sl.shapes.title, sl.placeholders[1]
            title.left, title.top, title.width, title.height = Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.4)
            body.left, body.top, body.width, body.height = Inches(0.9), Inches(2.1), Inches(11.5), Inches(4.8)
            title.text_frame.text = s["title"]
            title.text_frame.word_wrap = True
            title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            title.text_frame.paragraphs[0].runs[0].font.size = Pt(TITLE_PT)
            title.text_frame.paragraphs[0].runs[0].font.bold = True
            tf = body.text_frame
            tf.word_wrap = True
            for k, b in enumerate(s["bullets"]):
                par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                par.text = b
                par.runs[0].font.size = Pt(BODY_PT)
                par.space_after = Pt(10)
        sl.notes_slide.notes_text_frame.text = s["notes"]
    out = OUT_DIR / "slides.pptx"
    prs.save(out)
    print("saved", out, len(slides), "slides")


main()
