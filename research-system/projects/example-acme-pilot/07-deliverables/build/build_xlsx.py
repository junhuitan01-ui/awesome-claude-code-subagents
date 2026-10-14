import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from common import OUT_DIR, FORMATTED

SHEETS = [
    ("Monthly metrics", "monthly-metrics.csv",
     "Table 2: Monthly metrics, July 2025 to June 2026",
     "Note: Before = July to December 2025; Pilot = January to June 2026. Units shipped are units per month, "
     "complaints are a count per month, and sick days are days per month. The file does not define sick days, "
     "and working days per month may differ between periods.",
     "Source: [1] Acme Widgets Ltd monthly operations data (pilot-metrics.csv).",
     {"Units shipped", "On-time delivery (%)", "Customer complaints", "Sick days"}),
    ("Survey", "survey.csv",
     "Table 3: Staff wellbeing survey, December 2025 and June 2026",
     "Note: Before = December 2025 (41 of 50 responded, 82%). After = June 2026 (37 of 51 responded, reported as 73%). "
     "Five is the best score. The survey is anonymous, so respondents may differ between rounds.",
     "Source: [2] Acme Widgets Ltd staff wellbeing survey summary (staff-survey-summary.md).",
     {"Before", "After", "Change"}),
]


def conv(v):
    try:
        return int(v)
    except ValueError:
        pass
    try:
        return float(v)
    except ValueError:
        return v


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for name, fn, cap, note, src, numeric in SHEETS:
        ws = wb.create_sheet(name)
        with open(FORMATTED / "tables" / fn, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        header = rows[0]
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows[1:]:
            out = [conv(v) if header[k] in numeric else v for k, v in enumerate(r)]
            ws.append(out)
        # Survey "Change" is a signed text like +0.7; store as number
        # (conv handles it: float("+0.7") works).
        last = len(rows)
        for off, text, bold in ((2, cap, True), (3, note, False), (4, src, False)):
            cell = ws.cell(row=last + off, column=1, value=text)
            cell.font = Font(bold=bold, italic=not bold)
        for k, h in enumerate(header, start=1):
            width = max(len(str(h)), *(len(str(r[k - 1])) for r in rows[1:])) + 3
            ws.column_dimensions[ws.cell(row=1, column=k).column_letter].width = min(width, 50)
        ws.freeze_panes = "A2"
    out = OUT_DIR / "data.xlsx"
    wb.save(out)
    print("saved", out)


main()
