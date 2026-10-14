from docx import Document
from docx.shared import Pt
from common import OUT_DIR, parse_report, runs


def add_runs(par, text):
    for t, b, i, c in runs(text):
        r = par.add_run(t)
        r.bold = b or None
        r.italic = i or None
        if c:
            r.font.name = "Courier New"


def main():
    title, blocks = parse_report()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=0)
    prev_table = False
    for kind, p in blocks:
        if kind == "h2":
            doc.add_heading(p, level=1)
        elif kind == "h3":
            doc.add_heading(p, level=2)
        elif kind == "p":
            add_runs(doc.add_paragraph(), p)
        elif kind == "bullet":
            add_runs(doc.add_paragraph(style="List Bullet"), p)
        elif kind == "num":
            add_runs(doc.add_paragraph(style="List Number"), p)
        elif kind == "table":
            t = doc.add_table(rows=len(p), cols=len(p[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(p):
                for ci, cell in enumerate(row):
                    par = t.cell(ri, ci).paragraphs[0]
                    add_runs(par, cell)
                    for r in par.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0:
                            r.bold = True
            doc.add_paragraph()
    out = OUT_DIR / "report.docx"
    doc.save(out)
    print("saved", out)


main()
