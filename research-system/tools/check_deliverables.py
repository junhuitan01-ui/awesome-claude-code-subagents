#!/usr/bin/env python3
"""Check that a project's deliverables exist, open, and contain real content.

Usage:
    python3 tools/check_deliverables.py projects/<name>          # summary
    python3 tools/check_deliverables.py projects/<name> --dump   # also print extracted text

Expected files are read from the "## Files to deliver" section of
projects/<name>/00-brief/deliverable-spec.md (lines like "- `slides.pptx`").
Exit code is 0 when every expected file passes, 1 otherwise.

Only the Python standard library is required. Opening .docx/.pptx/.xlsx/.pdf
uses python-docx, python-pptx, openpyxl and pypdf when installed; if one is
missing, that file is reported as "not verified" rather than passed.
"""
import argparse
import re
import sys
import zipfile
from pathlib import Path

PLACEHOLDER_PATTERNS = [r"\[EVIDENCE GAP", r"\bTODO\b", r"\bTBD\b", r"lorem ipsum", r"\[[SI]\d+\]"]


def expected_files(project: Path):
    spec = project / "00-brief" / "deliverable-spec.md"
    if not spec.exists():
        return None
    names, in_section = [], False
    for line in spec.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            in_section = line.strip().lower().startswith("## files to deliver")
            continue
        m = re.match(r"\s*[-*]\s+`?([\w.\- ]+\.\w+)`?", line) if in_section else None
        if m:
            names.append(m.group(1).strip())
    return names


def read_docx(path):
    import docx
    d = docx.Document(str(path))
    heads = [p.text for p in d.paragraphs if p.style.name.lower().startswith("heading") and p.text.strip()]
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return f"{len(text.split())} words, {len(heads)} headings, {len(d.tables)} tables", text


def read_pptx(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    parts, notes = [], 0
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        parts.append(f"--- slide {i} ---\n" + "\n".join(texts))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes += 1
    return f"{len(prs.slides)} slides, {notes} with speaker notes", "\n".join(parts)


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    summary, parts = [], []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        summary.append(f"'{ws.title}': {len(rows)} rows")
        parts.append(f"--- sheet {ws.title} ---\n" + "\n".join(" | ".join("" if v is None else str(v) for v in r) for r in rows[:50]))
    return "; ".join(summary), "\n".join(parts)


def read_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(str(path))
    text = "\n".join((p.extract_text() or "") for p in r.pages)
    return f"{len(r.pages)} pages, {len(text.split())} words", text


def read_text(path):
    text = path.read_text(encoding="utf-8", errors="replace")
    return f"{len(text.splitlines())} lines, {len(text.split())} words", text


READERS = {".docx": read_docx, ".pptx": read_pptx, ".xlsx": read_xlsx, ".pdf": read_pdf,
           ".md": read_text, ".csv": read_text, ".txt": read_text, ".html": read_text}


def check(path: Path):
    """Return (status, details, text). status: PASS, FAIL or NOT VERIFIED."""
    if not path.exists():
        return "FAIL", "file missing", ""
    if path.stat().st_size == 0:
        return "FAIL", "file is empty", ""
    ext = path.suffix.lower()
    if ext in (".docx", ".pptx", ".xlsx") and not zipfile.is_zipfile(path):
        return "FAIL", "not a valid Office file", ""
    if ext == ".pdf" and not path.read_bytes()[:5] == b"%PDF-":
        return "FAIL", "not a valid PDF", ""
    reader = READERS.get(ext)
    if reader is None:
        return "NOT VERIFIED", f"no reader for {ext} files ({path.stat().st_size} bytes)", ""
    try:
        details, text = reader(path)
    except ImportError as e:
        return "NOT VERIFIED", f"library not installed to open it ({e.name})", ""
    except Exception as e:  # corrupt or unreadable file
        return "FAIL", f"could not open: {e}", ""
    if len(text.split()) < 20:
        return "FAIL", f"almost no content ({details})", text
    found = [p for p in PLACEHOLDER_PATTERNS if re.search(p, text, re.I)]
    if found:
        details += "; contains markers: " + ", ".join(found)
    return "PASS", details, text


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("project", help="path to projects/<name>")
    ap.add_argument("--dump", action="store_true", help="print the extracted text of each file")
    args = ap.parse_args()

    project = Path(args.project)
    out_dir = project / "07-deliverables"
    expected = expected_files(project)
    if expected is None:
        print(f"WARNING: {project}/00-brief/deliverable-spec.md not found; checking whatever is in {out_dir}")
        expected = []
    present = sorted(p.name for p in out_dir.glob("*") if p.is_file()) if out_dir.exists() else []
    names = expected + [n for n in present if n not in expected]
    if not names:
        print(f"FAIL: no deliverables listed in the spec and none found in {out_dir}")
        return 1

    failed = False
    print(f"Deliverables check for {project}\n")
    for name in names:
        status, details, text = check(out_dir / name)
        tag = "" if name in expected else " (not in spec)"
        print(f"[{status}] {name}{tag}: {details}")
        if name in expected and status != "PASS":
            failed = True
        if args.dump and text:
            print(text[:6000] + ("\n...[truncated]" if len(text) > 6000 else "") + "\n")
    print("\nResult:", "FAILED (see above)" if failed else "all expected files passed")
    print("Note: this checks that files exist, open and have content. "
          "Whether the content meets the brief is judged in the final validation step.")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
